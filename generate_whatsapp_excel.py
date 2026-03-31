import json
import os
from pathlib import Path

# Ensure openpyxl is installed
try:
    import openpyxl
except ImportError:
    import subprocess, sys
    subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl"])
    import openpyxl

# Define paths relative to this script
BASE_DIR = Path(__file__).resolve().parent
JSON_PATH = BASE_DIR / "participants_master.json"
OUTPUT_XLSX = BASE_DIR / "whatsapp_links.xlsx"

# Load participants data
with open(JSON_PATH, "r", encoding="utf-8") as f:
    data = json.load(f)

# Extract unique phone numbers (digits only, prepend India code if missing)
phones = set()
for team in data:
    for member in team.get("members", []):
        raw = member.get("phone", "")
        digits = "".join(filter(str.isdigit, str(raw)))
        if len(digits) >= 10:
            if not digits.startswith("91"):
                digits = "91" + digits
            phones.add(digits)

phones = sorted(phones)

# URL‑encoded message (same as provided earlier)
encoded_message = (
    "Dear%20Team%20Head%2C%0A%0AGreetings%20from%20the%20AIExplore%202K26%20Organizing%20Committee.%0A%0A"
    "We%20would%20like%20to%20inform%20you%20about%20important%20updates%20regarding%20the%20event%20structure%20and%20schedule.%0A%0A"
    "AIExplore%202K26%20has%20been%20postponed%20and%20will%20now%20be%20conducted%20later%20this%20month.%20The%20revised%20dates%20will%20be%20announced%20shortly%20on%20the%20official%20Unstop%20page.%20We%20request%20all%20teams%20to%20regularly%20check%20the%20platform%20for%20updates.%0A%0A"
    "Round%201%20(Online%20PPT%20Shortlisting%20Round)%20remains%20completely%20free%20for%20all%20participants.%20Only%20shortlisted%20teams%20will%20advance%20to%20Round%202.%0A%0A"
    "Round%202%20will%20be%20the%2024-hour%20offline%20hackathon.%20Participation%20in%20this%20final%20round%20will%20require%20a%20mandatory%20registration%20fee%20per%20team.%20The%20exact%20amount%20and%20payment%20details%20will%20be%20communicated%20only%20to%20shortlisted%20teams.%0A%0A"
    "We%20appreciate%20your%20cooperation%20and%20enthusiasm.%20Kindly%20ensure%20your%20team%20is%20prepared%20for%20the%20upcoming%20updates%20and%20submission%20deadlines.%0A%0A"
    "For%20any%20queries%2C%20please%20feel%20free%20to%20contact%3A%0A%0AEvent%20Heads%3A%0AAditya%20Thakre%20%E2%80%93%207058045278%0AShubham%20Sharnagate%20%E2%80%93%209359910468%0A%0ATechnical%20Head%3A%0AAryan%20Motghare%20%E2%80%93%207248955703%0A%0ABest%20regards%2C%0AAIExplore%202K26%20Organizing%20Committee"
)

# Create workbook
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "WhatsApp Links"

# Header
ws["A1"] = "Phone"
ws["B1"] = "Encoded Message"
ws["C1"] = "WhatsApp Link"
ws["B1"].value = encoded_message

# Populate rows
for i, phone in enumerate(phones, start=2):
    ws.cell(row=i, column=1, value=phone)
    # Hyperlink formula referencing B1 for the message
    formula = f"=HYPERLINK(\"https://wa.me/\"&A{i}&\"?text=\"&$B$1, \"Send WhatsApp\")"
    ws.cell(row=i, column=3, value=formula)

# Adjust column widths for readability
ws.column_dimensions["A"].width = 15
ws.column_dimensions["B"].width = 120
ws.column_dimensions["C"].width = 40

# Save workbook
wb.save(OUTPUT_XLSX)
print(f"Excel file generated at {OUTPUT_XLSX}")
